import csv
import io
import math
from datetime import date, datetime, time, timezone
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from fastapi.responses import StreamingResponse
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.core.auth import get_current_user
from app.core.role_checker import normalize_role, require_role
from app.database.database import get_db
from app.models.activity import Activity
from app.models.contract import Contract
from app.models.user import User
from app.schemas.activity_schema import ActivityPaginationResponse, ActivityResponse

router = APIRouter(
    tags=["Activity Logs"]
)


def apply_rbac_activity_scope(query, current_user: User, db: Session):
    """
    Enforces RBAC scoping on activity log views:
    - Admin: View all logs
    - Compliance Officer: View all logs
    - Manager (Legal Manager, Contract Manager): View department / managed contract logs & own activities
    - Employee (Viewer): View only their own activities
    """
    role = normalize_role(current_user.role)

    if role in ["Admin", "Compliance Officer"]:
        return query

    if role in ["Legal Manager", "Contract Manager"]:
        # Managers can see activities they created, or activities on contracts they own/manage/assigned
        managed_contract_ids = [
            c[0]
            for c in db.query(Contract.id)
            .filter(
                or_(
                    Contract.created_by == current_user.id,
                    Contract.assigned_to == current_user.id,
                )
            )
            .all()
        ]

        conditions = [Activity.user_id == current_user.id]
        if managed_contract_ids:
            conditions.append(Activity.contract_id.in_(managed_contract_ids))

        return query.filter(or_(*conditions))

    # Viewer / Employee: View only their own activities
    return query.filter(Activity.user_id == current_user.id)


def build_filtered_activity_query(
    db: Session,
    current_user: User,
    search: Optional[str] = None,
    action: Optional[str] = None,
    status_val: Optional[str] = None,
    user_id: Optional[int] = None,
    role: Optional[str] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
):
    query = db.query(Activity)
    query = apply_rbac_activity_scope(query, current_user, db)

    if search and search.strip():
        term = f"%{search.strip()}%"
        query = query.filter(
            or_(
                Activity.description.ilike(term),
                Activity.activity.ilike(term),
                Activity.user_name.ilike(term),
                Activity.action.ilike(term),
                Activity.entity_type.ilike(term),
                Activity.ip_address.ilike(term),
            )
        )

    if action and action.strip():
        query = query.filter(Activity.action == action.strip())

    if status_val and status_val.strip():
        query = query.filter(Activity.status.ilike(status_val.strip()))

    if user_id:
        query = query.filter(Activity.user_id == user_id)

    if role and role.strip():
        query = query.filter(Activity.user_role == normalize_role(role.strip()))

    if start_date:
        start_dt = datetime.combine(start_date, time.min, tzinfo=timezone.utc)
        query = query.filter(Activity.timestamp >= start_dt)

    if end_date:
        end_dt = datetime.combine(end_date, time.max, tzinfo=timezone.utc)
        query = query.filter(Activity.timestamp <= end_dt)

    return query


@router.get(
    "/activity",
    response_model=ActivityPaginationResponse,
    status_code=status.HTTP_200_OK,
)
@router.get(
    "/activities",
    response_model=ActivityPaginationResponse,
    status_code=status.HTTP_200_OK,
)
def get_activities(
    page: int = Query(1, ge=1, description="Page number"),
    page_size: int = Query(20, ge=1, le=100, description="Items per page"),
    sort_by: str = Query("timestamp", description="Sort field (timestamp, id, action)"),
    order: str = Query("desc", pattern="^(asc|desc)$", description="Sort direction"),
    search: Optional[str] = Query(None, description="Search text across description, user, action"),
    action: Optional[str] = Query(None, description="Filter by action code"),
    status: Optional[str] = Query(None, description="Filter by status (Success, Failed)"),
    user_id: Optional[int] = Query(None, description="Filter by user ID"),
    role: Optional[str] = Query(None, description="Filter by user role"),
    start_date: Optional[date] = Query(None, description="Filter by start date"),
    end_date: Optional[date] = Query(None, description="Filter by end date"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Retrieve paginated, filtered, and role-scoped activity audit logs.
    """
    query = build_filtered_activity_query(
        db=db,
        current_user=current_user,
        search=search,
        action=action,
        status_val=status,
        user_id=user_id,
        role=role,
        start_date=start_date,
        end_date=end_date,
    )

    total = query.count()

    # Apply sorting
    sort_col = Activity.timestamp
    if sort_by == "id":
        sort_col = Activity.id
    elif sort_by == "action":
        sort_col = Activity.action

    if order.lower() == "asc":
        query = query.order_by(sort_col.asc(), Activity.id.asc())
    else:
        query = query.order_by(sort_col.desc(), Activity.id.desc())

    offset = (page - 1) * page_size
    items = query.offset(offset).limit(page_size).all()
    total_pages = math.ceil(total / page_size) if total > 0 else 1

    return ActivityPaginationResponse(
        items=items,
        total=total,
        page=page,
        page_size=page_size,
        total_pages=total_pages,
    )


@router.get(
    "/activity/export/csv",
    status_code=status.HTTP_200_OK,
)
@router.get(
    "/activities/export/csv",
    status_code=status.HTTP_200_OK,
)
def export_activities_csv(
    search: Optional[str] = None,
    action: Optional[str] = None,
    status: Optional[str] = None,
    user_id: Optional[int] = None,
    role: Optional[str] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Export activity logs as a CSV spreadsheet.
    """
    query = build_filtered_activity_query(
        db=db,
        current_user=current_user,
        search=search,
        action=action,
        status_val=status,
        user_id=user_id,
        role=role,
        start_date=start_date,
        end_date=end_date,
    ).order_by(Activity.timestamp.desc()).limit(1000)

    records = query.all()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "ID", "Timestamp (UTC)", "User Name", "Role", "Action",
        "Entity Type", "Entity ID", "Contract ID", "Description", "IP Address", "Status"
    ])

    for r in records:
        ts_str = r.timestamp.strftime("%Y-%m-%d %H:%M:%S UTC") if r.timestamp else ""
        writer.writerow([
            r.id,
            ts_str,
            r.user_name or "System",
            r.user_role or "",
            r.action,
            r.entity_type or "",
            r.entity_id or "",
            r.contract_id or "",
            r.description or r.activity or "",
            r.ip_address or "",
            r.status or "Success",
        ])

    output.seek(0)
    filename = f"activity_logs_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.csv"
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8")),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get(
    "/activity/export/excel",
    status_code=status.HTTP_200_OK,
)
@router.get(
    "/activities/export/excel",
    status_code=status.HTTP_200_OK,
)
def export_activities_excel(
    search: Optional[str] = None,
    action: Optional[str] = None,
    status: Optional[str] = None,
    user_id: Optional[int] = None,
    role: Optional[str] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Export activity logs as an Excel spreadsheet (.xlsx).
    """
    query = build_filtered_activity_query(
        db=db,
        current_user=current_user,
        search=search,
        action=action,
        status_val=status,
        user_id=user_id,
        role=role,
        start_date=start_date,
        end_date=end_date,
    ).order_by(Activity.timestamp.desc()).limit(1000)

    records = query.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Activity Logs"
    ws.views.sheetView[0].showGridLines = True

    # Banner
    ws.merge_cells("A1:K2")
    banner = ws.cell(row=1, column=1, value="ContractIQ — Operational Activity Logs Export")
    banner.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    banner.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    banner.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # Headers
    headers = [
        "ID", "Timestamp (UTC)", "User Name", "Role", "Action",
        "Entity Type", "Entity ID", "Contract ID", "Description", "IP Address", "Status"
    ]
    th_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    th_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = th_font
        cell.fill = th_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    data_font = Font(name="Calibri", size=10, color="1E293B")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for row_idx, r in enumerate(records, start=5):
        ts_str = r.timestamp.strftime("%Y-%m-%d %H:%M:%S UTC") if r.timestamp else ""
        vals = [
            r.id,
            ts_str,
            r.user_name or "System",
            r.user_role or "",
            r.action,
            r.entity_type or "",
            r.entity_id or "",
            r.contract_id or "",
            r.description or r.activity or "",
            r.ip_address or "",
            r.status or "Success",
        ]
        for col_idx, val in enumerate(vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in [1, 2]:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"activity_logs_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get(
    "/activity/export/pdf",
    status_code=status.HTTP_200_OK,
)
@router.get(
    "/activities/export/pdf",
    status_code=status.HTTP_200_OK,
)
def export_activities_pdf(
    search: Optional[str] = None,
    action: Optional[str] = None,
    status: Optional[str] = None,
    user_id: Optional[int] = None,
    role: Optional[str] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Export activity logs as a PDF document.
    """
    query = build_filtered_activity_query(
        db=db,
        current_user=current_user,
        search=search,
        action=action,
        status_val=status,
        user_id=user_id,
        role=role,
        start_date=start_date,
        end_date=end_date,
    ).order_by(Activity.timestamp.desc()).limit(200)

    records = query.all()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "PDFTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0F172A"),
    )
    th_style = ParagraphStyle(
        "PDFTH",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white,
    )
    td_style = ParagraphStyle(
        "PDFTD",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7,
        leading=9,
        textColor=colors.HexColor("#1E293B"),
    )

    elements = []
    elements.append(Paragraph("<b>ContractIQ</b> — Operational Activity Logs", title_style))
    elements.append(Spacer(1, 4))
    elements.append(Paragraph(f"Export generated on: {datetime.now(timezone.utc).strftime('%B %d, %Y - %H:%M UTC')} | Requested by: {current_user.full_name} ({current_user.role})", ParagraphStyle("Sub", fontName="Helvetica", fontSize=8.5, textColor=colors.HexColor("#64748B"))))
    elements.append(Spacer(1, 10))

    table_data = [[
        Paragraph("ID", th_style),
        Paragraph("Timestamp", th_style),
        Paragraph("User", th_style),
        Paragraph("Role", th_style),
        Paragraph("Action", th_style),
        Paragraph("Description", th_style),
        Paragraph("Status", th_style),
    ]]

    for r in records:
        ts_str = r.timestamp.strftime("%Y-%m-%d %H:%M") if r.timestamp else ""
        table_data.append([
            Paragraph(str(r.id), td_style),
            Paragraph(ts_str, td_style),
            Paragraph(r.user_name or "System", td_style),
            Paragraph(r.user_role or "—", td_style),
            Paragraph(r.action, td_style),
            Paragraph(r.description or r.activity or "", td_style),
            Paragraph(r.status or "Success", td_style),
        ])

    tbl = Table(table_data, colWidths=[30, 80, 80, 65, 90, 150, 45], repeatRows=1)
    t_styles = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for i in range(1, len(table_data)):
        if i % 2 == 0:
            t_styles.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F8FAFC")))
    tbl.setStyle(TableStyle(t_styles))
    elements.append(tbl)

    def draw_footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(36, 20, "ContractIQ Activity Audit Trail — Confidential")
        canvas.drawRightString(document.pagesize[0] - 36, 20, f"Page {document.page}")
        canvas.restoreState()

    doc.build(elements, onFirstPage=draw_footer, onLaterPages=draw_footer)
    buf.seek(0)
    filename = f"activity_logs_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get(
    "/activity/{activity_id}",
    response_model=ActivityResponse,
    status_code=status.HTTP_200_OK,
)
@router.get(
    "/activities/{activity_id}",
    response_model=ActivityResponse,
    status_code=status.HTTP_200_OK,
)
def get_activity(
    activity_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Retrieve single activity details with permission checks.
    """
    query = db.query(Activity).filter(Activity.id == activity_id)
    query = apply_rbac_activity_scope(query, current_user, db)
    activity = query.first()

    if not activity:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Activity log not found",
        )

    return activity


@router.delete(
    "/activity/{activity_id}",
    status_code=status.HTTP_200_OK,
)
@router.delete(
    "/activities/{activity_id}",
    status_code=status.HTTP_200_OK,
)
def delete_activity(
    activity_id: int,
    current_user: User = Depends(require_role("Admin")),
    db: Session = Depends(get_db),
):
    """
    Delete an activity audit log. Restricted to Administrator role only.
    """
    activity = db.query(Activity).filter(Activity.id == activity_id).first()
    if not activity:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Activity log not found",
        )

    db.delete(activity)
    db.commit()

    return {
        "detail": f"Activity {activity_id} deleted successfully",
        "id": activity_id,
    }